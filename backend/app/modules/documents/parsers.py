"""Document parsers for PDF, Excel/CSV, and images."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from typing import Protocol

from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.config import settings


@dataclass
class ParsedBlock:
    text: str
    metadata: dict = field(default_factory=dict)


@dataclass
class ParseResult:
    blocks: list[ParsedBlock]
    metadata: dict = field(default_factory=dict)


class DocumentParser(Protocol):
    def parse(self, data: bytes, *, file_name: str = "") -> ParseResult: ...


class PdfTextParser:
    def parse(self, data: bytes, *, file_name: str = "") -> ParseResult:
        reader = PdfReader(io.BytesIO(data))
        blocks: list[ParsedBlock] = []
        for idx, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                blocks.append(ParsedBlock(text=text, metadata={"page": idx, "file_name": file_name}))
        return ParseResult(blocks=blocks, metadata={"parser": "PdfTextParser", "pages": len(reader.pages)})


class ExcelCsvParser:
    def parse(self, data: bytes, *, file_name: str = "") -> ParseResult:
        name = file_name.lower()
        if name.endswith(".csv") or (not name.endswith((".xlsx", ".xls")) and b"," in data[:200]):
            return self._parse_csv(data, file_name=file_name)
        return self._parse_excel(data, file_name=file_name)

    def _parse_csv(self, data: bytes, *, file_name: str) -> ParseResult:
        text = data.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        blocks: list[ParsedBlock] = []
        for i in range(0, len(rows), 20):
            chunk_rows = rows[i : i + 20]
            lines = [", ".join(cell.strip() for cell in row if str(cell).strip()) for row in chunk_rows]
            lines = [line for line in lines if line]
            if lines:
                blocks.append(
                    ParsedBlock(
                        text="\n".join(lines),
                        metadata={"sheet": "csv", "row_start": i + 1, "row_end": i + len(chunk_rows), "file_name": file_name},
                    )
                )
        return ParseResult(blocks=blocks, metadata={"parser": "ExcelCsvParser", "format": "csv", "rows": len(rows)})

    def _parse_excel(self, data: bytes, *, file_name: str) -> ParseResult:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        blocks: list[ParsedBlock] = []
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            for i in range(0, len(rows), 20):
                chunk_rows = rows[i : i + 20]
                lines = []
                for row in chunk_rows:
                    cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if cells:
                        lines.append(" | ".join(cells))
                if lines:
                    blocks.append(
                        ParsedBlock(
                            text="\n".join(lines),
                            metadata={
                                "sheet": sheet.title,
                                "row_start": i + 1,
                                "row_end": i + len(chunk_rows),
                                "file_name": file_name,
                            },
                        )
                    )
        return ParseResult(blocks=blocks, metadata={"parser": "ExcelCsvParser", "format": "excel", "sheets": len(wb.sheetnames)})


class ImageVisionParser:
    """Uses OpenAI vision when configured; otherwise returns a placeholder description."""

    def parse(self, data: bytes, *, file_name: str = "") -> ParseResult:
        if settings.openai_api_key:
            try:
                from openai import OpenAI

                import base64

                client = OpenAI(api_key=settings.openai_api_key)
                b64 = base64.b64encode(data).decode("ascii")
                ext = "png" if file_name.lower().endswith(".png") else "jpeg"
                response = client.chat.completions.create(
                    model=settings.openai_model,
                    messages=[
                        {
                            "role": "user",
                            "content": [
                                {
                                    "type": "text",
                                    "text": (
                                        "Extract construction-relevant details from this drawing/image: "
                                        "dimensions, materials, quantities, room labels, and specs. "
                                        "Return plain text."
                                    ),
                                },
                                {"type": "image_url", "image_url": {"url": f"data:image/{ext};base64,{b64}"}},
                            ],
                        }
                    ],
                    max_tokens=1200,
                )
                text = (response.choices[0].message.content or "").strip()
                return ParseResult(
                    blocks=[ParsedBlock(text=text, metadata={"file_name": file_name, "source": "vision"})],
                    metadata={"parser": "ImageVisionParser", "mode": "openai"},
                )
            except Exception as exc:  # noqa: BLE001
                text = f"[Vision parse failed: {exc}] Image upload: {file_name}"
        else:
            text = (
                f"[Image document: {file_name}] No OPENAI_API_KEY configured. "
                "Vision extraction skipped; upload a PDF/Excel BOQ for text extraction."
            )
        return ParseResult(
            blocks=[ParsedBlock(text=text, metadata={"file_name": file_name, "source": "fallback"})],
            metadata={"parser": "ImageVisionParser", "mode": "fallback"},
        )


class ParserFactory:
    def for_type(self, file_type: str) -> DocumentParser:
        if file_type == "pdf":
            return PdfTextParser()
        if file_type in {"excel", "csv"}:
            return ExcelCsvParser()
        if file_type == "image":
            return ImageVisionParser()
        raise ValueError(f"Unsupported file type: {file_type}")
